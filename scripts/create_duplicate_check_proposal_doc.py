from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path("docs/duplicate_check_proposal.docx")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)

    for name, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def style_run(run, bold=False, size=11, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_title(doc, text, subtitle):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    style_run(r, bold=True, size=20, color="0B2545")

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run(subtitle)
    style_run(r, size=12, color="555555")


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.style = f"Heading {level}"
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    style_run(r, bold=True, size=16 if level == 1 else 13, color="2E74B5")


def add_body(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.10
    r = p.add_run(text)
    style_run(r)


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.167
    r = p.add_run(text)
    style_run(r)


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.167
    r = p.add_run(text)
    style_run(r)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.width = widths[idx]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "F2F4F7")
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        style_run(r, bold=True)

    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cell = cells[idx]
            cell.width = widths[idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.10
            r = p.add_run(value)
            style_run(r)

    doc.add_paragraph()


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(11)
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    add_title(
        doc,
        "Duplicate Registration Check Proposal",
        "ข้อเสนอ logic ใหม่สำหรับกันสมัครซ้ำในรอบเปิดรับปัจจุบัน",
    )

    add_body(
        doc,
        "เอกสารนี้สรุปแนวทางที่ตกลงกันในเชิง requirement ยังไม่ได้หมายความว่าโค้ดถูกปรับตาม logic นี้แล้ว",
    )

    add_heading(doc, "ขอบเขตที่ต้องเช็ค", 1)
    add_body(
        doc,
        "ให้เช็คเฉพาะข้อมูลสมัครในรอบเปิดรับปัจจุบันของสนาม/ฟอร์มนั้น โดยใช้ opening_time และ closing_time ล่าสุดของ route/form นั้นเป็นกรอบเวลา",
    )
    add_bullet(doc, "ไม่เช็คข้ามรอบเปิดรับ เพราะลูกค้าคนเดิมควรสมัครรอบใหม่ได้")
    add_bullet(doc, "ไม่เช็คข้ามทุกสนามหรือทุกฟอร์ม เพราะลูกค้าอาจสมัครหลายรายการได้")
    add_bullet(doc, "เช็คเฉพาะ record ใน receipt_forms ที่มี submission_date อยู่ในรอบเปิดรับปัจจุบัน")

    add_heading(doc, "ลำดับการเช็ค", 1)
    add_number(doc, "เมื่อกรอกชื่อและนามสกุลครบ ให้เช็คชื่อ + นามสกุลก่อนทันที")
    add_number(doc, "ถ้าชื่อ + นามสกุลซ้ำในรอบเปิดรับปัจจุบัน ให้แจ้งเตือนตรง field ชื่อ/นามสกุล และกันการ submit")
    add_number(doc, "ถ้าชื่อ + นามสกุลไม่ซ้ำ จึงค่อยเช็ค email เป็นลำดับถัดไป")
    add_number(doc, "ถ้า email ซ้ำในรอบเปิดรับปัจจุบัน ให้แจ้งเตือนตรง field email และกันการ submit")
    add_number(doc, "ตอน submit ต้องเช็คซ้ำฝั่ง backend อีกครั้งด้วยลำดับเดียวกัน เพื่อกันการยิง request ตรงหรือ race condition")

    add_heading(doc, "เงื่อนไขการ Normalize", 1)
    add_table(
        doc,
        ["ข้อมูล", "วิธี normalize"],
        [
            ["ชื่อ / นามสกุล", "trim, รวมช่องว่างซ้ำ, แปลงเป็น lowercase"],
            ["Email", "trim และแปลงเป็น lowercase"],
            ["รอบเปิดรับ", "ใช้ opening_time ถึง closing_time ล่าสุดของ route/form"],
        ],
        [Inches(1.7), Inches(4.6)],
    )

    add_heading(doc, "พฤติกรรมที่ต้องการบนหน้า Form", 1)
    add_table(
        doc,
        ["จังหวะ", "การทำงาน"],
        [
            ["กรอกชื่อ/นามสกุลครบ", "เรียก endpoint เพื่อตรวจชื่อซ้ำในรอบเปิดรับปัจจุบัน"],
            ["ชื่อ/นามสกุลซ้ำ", "แสดง error ใต้ field ชื่อหรือนามสกุลทันที"],
            ["ชื่อ/นามสกุลไม่ซ้ำ", "ปล่อยให้กรอกต่อ และเมื่อกรอก email ให้ตรวจ email ซ้ำ"],
            ["Email ซ้ำ", "แสดง error ใต้ field email ทันที"],
            ["กด Submit", "backend ตรวจซ้ำอีกรอบตามลำดับเดิมก่อนบันทึก"],
        ],
        [Inches(2.0), Inches(4.3)],
    )

    add_heading(doc, "เหตุผลของแนวทางนี้", 1)
    add_bullet(doc, "กันเคสลูกค้าเปลี่ยนอีเมลเพื่อสมัครซ้ำได้ตั้งแต่ขั้นชื่อ/นามสกุล")
    add_bullet(doc, "แจ้งเตือนเร็วขึ้น ไม่ต้องรอ submit ทั้งฟอร์ม")
    add_bullet(doc, "ยังเปิดให้ลูกค้าเก่ากลับมาสมัครรอบใหม่ได้ เพราะเช็คเฉพาะรอบเปิดรับ")
    add_bullet(doc, "ลดภาระแอดมินในการตรวจ duplicate หลังบ้าน")

    add_heading(doc, "ข้อควรระวัง", 1)
    add_bullet(doc, "ชื่อ + นามสกุลอย่างเดียวมีโอกาสชนคนละคนได้ แต่ตาม requirement นี้ต้องการห้ามซ้ำทันที")
    add_bullet(doc, "ข้อความเตือนควรบอกให้ติดต่อแอดมิน หากคิดว่าเป็นคนละคนแต่ชื่อซ้ำ")
    add_bullet(doc, "ควรมี backend validation เสมอ แม้มี AJAX validation หน้าเว็บแล้ว")

    doc.save(OUTPUT)


if __name__ == "__main__":
    main()
