from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path("docs/duplicate_registration_logic.docx")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)

    for margin_name, margin_value in {
        "top": top,
        "start": start,
        "bottom": bottom,
        "end": end,
    }.items():
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(margin_value))
        node.set(qn("w:type"), "dxa")


def style_run(run, bold=False, size=11, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph()
    paragraph.style = f"Heading {level}"
    run = paragraph.add_run(text)
    style_run(run, bold=True, size=16 if level == 1 else 13, color="2E74B5")
    return paragraph


def add_body(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.10
    run = paragraph.add_run(text)
    style_run(run)
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.167
    run = paragraph.add_run(text)
    style_run(run)
    return paragraph


def add_number(doc, text):
    paragraph = doc.add_paragraph(style="List Number")
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.167
    run = paragraph.add_run(text)
    style_run(run)
    return paragraph


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"

    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.width = widths[idx]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "F2F4F7")
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(header)
        style_run(run, bold=True)

    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cell = cells[idx]
            cell.width = widths[idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.10
            run = paragraph.add_run(value)
            style_run(run)

    doc.add_paragraph()
    return table


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(11)
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(3)
    run = title.add_run("Duplicate Registration Logic")
    style_run(run, bold=True, size=20, color="0B2545")

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(12)
    run = subtitle.add_run("ระบบกันสมัครซ้ำสำหรับฟอร์มสมัคร")
    style_run(run, size=12, color="555555")

    add_body(
        doc,
        "เอกสารนี้สรุป logic ปัจจุบันที่ใช้กันข้อมูลสมัครซ้ำในฟอร์ม dynamic "
        "หลังจากปรับจากการเช็คอีเมล มาเป็นการเช็คตัวตนผู้สมัครในรอบเปิดรับเดียวกัน",
    )

    add_heading(doc, "เป้าหมาย", 1)
    add_bullet(doc, "กันกรณีผู้สมัครคนเดิมเปลี่ยนอีเมลแล้วสมัครเข้ามาซ้ำ")
    add_bullet(doc, "ยังลดความเสี่ยงการ block คนละคนที่ชื่อและนามสกุลเหมือนกัน")
    add_bullet(doc, "คงขอบเขตการเช็คไว้เฉพาะฟอร์มเดียวกันและรอบเปิดรับเดียวกัน")

    add_heading(doc, "ขอบเขตการเช็ค", 1)
    add_body(
        doc,
        "ระบบจะเริ่มเช็คจากข้อมูลใน table receipt_forms โดยจำกัดเฉพาะ sfield ของฟอร์มนั้น "
        "และ submission_date ที่อยู่ระหว่าง opening_time ถึง closing_time ล่าสุดของ route/form นั้น",
    )

    add_table(
        doc,
        ["เงื่อนไข", "ความหมาย"],
        [
            ["ฟอร์มเดียวกัน", "ใช้ sfield เป็นตัวระบุฟอร์ม"],
            ["รอบเปิดรับเดียวกัน", "ใช้ opening_time และ closing_time ล่าสุดของ route/form"],
            ["ข้อมูลเดิม", "เช็คเฉพาะ record ใน receipt_forms ที่อยู่ในรอบนั้น"],
        ],
        [Inches(1.8), Inches(4.5)],
    )

    add_heading(doc, "การ Normalize ก่อนเช็ค", 1)
    add_table(
        doc,
        ["ข้อมูล", "วิธี normalize"],
        [
            ["ชื่อ / นามสกุล", "trim, รวมช่องว่างซ้ำให้เหลือหนึ่งช่อง, แปลงเป็น lowercase"],
            ["วันเกิด", "พยายามแปลงเป็นรูปแบบ Y-m-d"],
            ["เบอร์โทร", "ลบทุกอย่างที่ไม่ใช่ตัวเลข เช่น 081-234-5678 เป็น 0812345678"],
            ["อีเมล", "ไม่ใช้เป็นเงื่อนไขกันซ้ำแล้ว"],
        ],
        [Inches(1.8), Inches(4.5)],
    )

    add_heading(doc, "Logic หลัก", 1)
    add_number(doc, "ถ้ามีวันเกิด: ชื่อ + นามสกุล + วันเกิด ตรงกันในฟอร์มและรอบเดียวกัน = ซ้ำ")
    add_number(doc, "ถ้าไม่มีวันเกิด แต่มีเบอร์โทร: ชื่อ + นามสกุล + เบอร์โทร ตรงกันในฟอร์มและรอบเดียวกัน = ซ้ำ")
    add_number(doc, "ถ้าไม่มีทั้งวันเกิดและเบอร์โทร: ระบบจะไม่ถือว่าเป็นข้อมูลซ้ำ")
    add_number(doc, "อีเมลที่ต่างกันจะไม่ทำให้หลุด duplicate อีก ถ้าชื่อ + นามสกุล + วันเกิดตรงกัน")

    add_heading(doc, "ตัวอย่างผลลัพธ์", 1)
    add_table(
        doc,
        ["กรณี", "ผลลัพธ์"],
        [
            ["ชื่อ/นามสกุล/วันเกิดตรงกัน แต่อีเมลต่างกัน", "ถูก block เป็นข้อมูลซ้ำ"],
            ["ชื่อ/นามสกุลตรงกัน แต่วันเกิดต่างกัน", "ไม่ถือว่าซ้ำ"],
            ["ไม่มีวันเกิด แต่ชื่อ/นามสกุล/เบอร์ตรงกัน", "ถูก block เป็นข้อมูลซ้ำ"],
            ["สมัครรอบเก่า แล้วฟอร์มปิด จากนั้นเปิดรอบใหม่", "ไม่ถูกนับเป็นซ้ำ เพราะอยู่นอกรอบเปิดรับใหม่"],
        ],
        [Inches(3.1), Inches(3.2)],
    )

    add_heading(doc, "หมายเหตุสำคัญ", 1)
    add_bullet(
        doc,
        "การใช้เบอร์โทรเป็น fallback เท่านั้นช่วยลดปัญหา block พี่น้องหรือเด็กหลายคนที่ใช้เบอร์ผู้ปกครองเดียวกัน",
    )
    add_bullet(
        doc,
        "หากต้องการเข้มกว่านี้ในอนาคต สามารถเพิ่มเงื่อนไขอื่น เช่น เลขบัตร หรือรหัสผู้สมัครเฉพาะได้",
    )

    doc.save(OUTPUT)


if __name__ == "__main__":
    main()
